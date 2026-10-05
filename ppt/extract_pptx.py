import sys
from pptx import Presentation

path = sys.argv[1]
p = Presentation(path)
for i, slide in enumerate(p.slides):
    print(f'--- Slide {i+1} ---')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = ''.join(run.text for run in para.runs)
                if text.strip():
                    print(text)
        if shape.has_table:
            tbl = shape.table
            for row in tbl.rows:
                print(' | '.join(cell.text for cell in row.cells))
